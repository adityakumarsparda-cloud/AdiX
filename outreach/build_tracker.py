from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

NAVY = '1F3D5C'; F = 'Arial'
hdr_fill = PatternFill('solid', fgColor=NAVY)
hdr_font = Font(name=F, bold=True, color='FFFFFF', size=10)
warm_fill = PatternFill('solid', fgColor='FFF2CC')
thin = Side(style='thin', color='D0D5DA')
border = Border(left=thin, right=thin, top=thin, bottom=thin)

wb = Workbook()

# ---------------- Outreach sheet ----------------
ws = wb.active; ws.title = 'Outreach'
COLS = ['Date','Company','Website','Industry','Company Stage','Location','Contact Name',
        'Contact Designation','Contact Email','Contact Source','Warm/Cold','Target Role',
        'Email Subject','Email Sent','Date Sent','CV Attached','Response','Response Date',
        'Follow-up Date','Follow-up Status','Interview','Interview Date','Application Status','Notes']
WIDTHS = [11,20,26,20,14,14,17,22,32,26,10,24,38,11,11,12,10,13,13,16,10,13,17,52]

ws.append(COLS)
for i, w in enumerate(COLS, 1):
    c = ws.cell(1, i); c.fill = hdr_fill; c.font = hdr_font
    c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    ws.column_dimensions[get_column_letter(i)].width = WIDTHS[i-1]
ws.row_dimensions[1].height = 30
ws.freeze_panes = 'C2'

# Contacts drawn from Aditya's own Gmail history - every address below appeared in a real
# message to or from him, so none are guessed.
ROWS = [
 # date, company, website, industry, stage, location, name, designation, email, source, warm/cold,
 # role, subject, sent, datesent, cv, response, respdate, fudate, fustatus, interview, intdate, status, notes
 ['','MoreTasks','moretasks.com','Business Services / Ops','Growth','Delhi NCR','Gauri Chaturvedi',
  'Talent Acquisition','gauri.chaturvedi@moretasks.com','Gmail - inbound May 2026','Warm',
  'Operations Manager','','No','','','','','','Not started','No','','Ready to Send',
  'P1. Chased twice for Business Ops Associate, left callback 9205888799. Also cc: priyanshi.agarwal@moretasks.com, somya.devvanshi@moretasks.com'],
 ['','C11 Aura','','Consumer / Retail','Early','Kharar, Mohali','Navjot Kaur',
  'Recruiter','navjotkaurwork123@gmail.com','Gmail - inbound Jul 2026','Warm',
  'Operations Lead','','No','','','','','','Not started','No','','Ready to Send',
  'P1. Shortlisted for Ops Lead Round 2; interview missed (family emergency), they rescheduled 28 Jul. Closest role match to target.'],
 ['','The Bijou Box','thebijoubox.in','D2C / Jewellery','Early','Mumbai','Hiring Team',
  'Founder / Hiring','info@thebijoubox.in','Gmail - inbound May 2026','Warm',
  'E-commerce Operations','','No','','','','','','Not started','No','','Ready to Send',
  'P1. Said "CV aligns well with our requirements", booked a call, then went quiet. Requires Mumbai relocation.'],
 ['','Gudz','gudz.in','D2C / Consumer','Early','Gurugram (Sector 49)','Aryaman',
  'Founder','aryaman@gudz.in','Gmail - inbound May 2026','Warm',
  'Operations Manager','','No','','','','','','Not started','No','','Ready to Send',
  'P1. Replied same day, offered walk-in. Earlier role was fresher-level; re-approach at manager level. cc: darshan@gudz.in'],
 ['','Meesho','meesho.com','E-commerce Marketplace','Established','Panipat / NCR','Sharukh Khan',
  'Recruiter','sharukh.khan1@meesho.com','Gmail - Apr/May 2026','Warm',
  'Area Manager / Ops','','No','','','','','','Not started','No','','Ready to Send',
  'P1. Area Manager process was live via mynexthire; asked for update, no reply. Worth reopening.'],
 ['','Elda','getelda.com','Consumer Internet','Early','','Jayant',
  'Founder','jayant@getelda.com','Gmail - inbound Jul 2026','Warm',
  'Business Operations','','No','','','','','','Not started','No','','Ready to Send',
  "P1. Sent Founder's Office JD in Jul 2026. Was intern-level; ask about full-time ops."],
 ['','Cameron SF','cameronsf.com','D2C / Consumer','Early','','Justin',
  'Founder','justin@cameronsf.com','Gmail - outbound May 2026','Cold',
  'Operations Manager','','No','','','','','','Not started','No','','Ready to Send',
  'P3. Applied for Operations Manager May 2026, no reply. Verify company still active before re-sending.'],
 ['','Zoomuv','zoomuv.com','Logistics / Mobility','Early','','Dheeraj Gupta',
  'Founder','dheeraj@zoomuv.com','Gmail - outbound May 2026','Cold',
  'Operations Executive','','No','','','','','','Not started','No','','Ready to Send',
  'P3. Applied May 2026, no reply.'],
 ['','Ugaoo','ugaoo.com','D2C / Plants','Growth','Pune','Shubham S',
  'Hiring Manager','shubham.s@ugaoo.com','Gmail - outbound May 2026','Cold',
  'Operations Manager','','No','','','','','','Not started','No','','Ready to Send',
  'P2. Applied for Manager - Process Excellence May 2026. Established D2C, strong fit.'],
 ['','Samplify','gosamplify.com','D2C / Sampling','Early','','Hemlata',
  'HR','hemlata@gosamplify.com','Gmail - outbound May 2026','Cold',
  'Team Lead - Operations','','No','','','','','','Not started','No','','Ready to Send',
  'P2. Applied May 2026, no reply.'],
 ['','goPortals','goportals.co','E-commerce','Early','','Sahil Gupta',
  'Founder','sahilgupta@goportals.co','Gmail - outbound May 2026','Cold',
  'E-commerce Operations','','No','','','','','','Not started','No','','Ready to Send','P3. Applied May 2026, no reply.'],
 ['','CitizenEarth','citizenearth.in','D2C / Sustainability','Early','','Hiring Team',
  'HR','hr@citizenearth.in','Gmail - outbound May 2026','Cold',
  'Operations Lead','','No','','','','','','Not started','No','','Ready to Send','P4. Generic HR alias. Applied May 2026, no reply.'],
 ['','iCARRY','icarry.com','Logistics / E-commerce','Growth','','Talent Team',
  'Talent','Talents@icarry.com','Gmail - outbound May 2026','Cold',
  'Operations Manager','','No','','','','','','Not started','No','','Ready to Send','P4. Generic alias. Logistics OS for emerging markets - good ops fit.'],
]
for r in ROWS: ws.append(r)

for row in ws.iter_rows(min_row=2, max_row=ws.max_row, max_col=len(COLS)):
    for c in row:
        c.font = Font(name=F, size=10); c.border = border
        c.alignment = Alignment(vertical='top', wrap_text=(c.column in (24,)))
    if row[10].value == 'Warm':
        for c in row: c.fill = warm_fill

# Dropdowns
def dv(formula, col, last=400):
    d = DataValidation(type='list', formula1=formula, allow_blank=True)
    ws.add_data_validation(d); d.add(f'{col}2:{col}{last}')

dv('"Researching,Ready to Send,Sent,Replied,Follow-up Required,Interview,Rejected,No Response,Closed"','W')
dv('"Warm,Cold"','K')
dv('"Yes,No"','N'); dv('"Yes,No"','P'); dv('"Yes,No"','Q'); dv('"Yes,No"','R')
dv('"Not started,Due,Sent,Not needed"','T')
dv('"Early,Growth,Established"','E')

# ---------------- Dashboard ----------------
d = wb.create_sheet('Dashboard')
d.column_dimensions['A'].width = 34; d.column_dimensions['B'].width = 14
d['A1'] = 'Outreach Dashboard'; d['A1'].font = Font(name=F, bold=True, size=14, color=NAVY)
d['A2'] = 'Counts refresh when the file is opened in Excel or Google Sheets.'
d['A2'].font = Font(name=F, size=9, italic=True, color='666666')
METRICS = [
 ('Total rows in tracker', '=COUNTA(Outreach!B2:B400)'),
 ('Emails sent', '=COUNTIF(Outreach!N2:N400,"Yes")'),
 ('Warm contacts', '=COUNTIF(Outreach!K2:K400,"Warm")'),
 ('Cold contacts', '=COUNTIF(Outreach!K2:K400,"Cold")'),
 ('Responses received', '=COUNTIF(Outreach!Q2:Q400,"Yes")'),
 ('Interviews', '=COUNTIF(Outreach!U2:U400,"Yes")'),
 ('Follow-ups due', '=COUNTIF(Outreach!T2:T400,"Due")'),
 ('Ready to send', '=COUNTIF(Outreach!W2:W400,"Ready to Send")'),
 ('No response', '=COUNTIF(Outreach!W2:W400,"No Response")'),
 ('Response rate', '=IFERROR(COUNTIF(Outreach!Q2:Q400,"Yes")/COUNTIF(Outreach!N2:N400,"Yes"),0)'),
]
for i,(label,f) in enumerate(METRICS, start=4):
    d.cell(i,1,label).font = Font(name=F, size=10)
    c = d.cell(i,2,f); c.font = Font(name=F, size=10, bold=True)
    if label == 'Response rate': c.number_format = '0.0%'

d['A16'] = 'Daily cap'; d['A16'].font = Font(name=F, size=10)
d['B16'] = 20; d['B16'].font = Font(name=F, size=10, bold=True, color='0000FF')
d['C16'] = 'Blue = edit this. Cap set by Aditya; Gmail free tier allows ~500 recipients/day.'
d['C16'].font = Font(name=F, size=9, italic=True, color='666666')

# ---------------- Do Not Contact ----------------
x = wb.create_sheet('Do Not Contact')
x.append(['Email','Company','Reason'])
for i in range(1,4):
    c = x.cell(1,i); c.fill = hdr_fill; c.font = hdr_font
for w,col in zip([34,22,74],'ABC'): x.column_dimensions[col].width = w
for r in [
 ['reshma@techslash.com','Techslash','Likely bait. "Your TCS application is incomplete" via a redirect tracker; TCS does not recruit this way.'],
 ['neha@advert.co.in','Advert.co.in','Likely bait. "Accenture is hiring" sent from an advertising domain.'],
 ['eric@jobrightai.com','JobRight AI','Automated US job blasts, irrelevant to India ops roles.'],
 ['johnbutler@thesearchlogixgroup.org','SearchLogix Group','Caution. Asked salary expectations then "before moving forward". Legitimate recruiters never charge candidates.'],
]:
    x.append(r)
for row in x.iter_rows(min_row=2, max_row=x.max_row, max_col=3):
    for c in row:
        c.font = Font(name=F, size=10); c.alignment = Alignment(vertical='top', wrap_text=True); c.border = border

# ---------------- Legend ----------------
g = wb.create_sheet('How to use')
g.column_dimensions['A'].width = 100
LINES = [
 ('Outreach tracker - how to use', True),
 ('', False),
 ('One row per contact, never per company - two people at one company are two rows.', False),
 ('', False),
 ('Fill in as you go:', True),
 ('  Email Sent / Date Sent   - set the moment it actually goes out, not when drafted.', False),
 ('  Response / Response Date - Yes only for a human reply, not an autoresponder.', False),
 ('  Follow-up Date           - 5 to 7 working days after sending. Most replies come from the second email.', False),
 ('  Application Status       - Researching, Ready to Send, Sent, Replied, Follow-up Required,', False),
 ('                             Interview, Rejected, No Response, Closed.', False),
 ('', False),
 ('Priority order (in Notes as P1-P4):', True),
 ('  P1  Existing warm contact - someone who already replied to Aditya.', False),
 ('  P2  Verified HR, recruiter or hiring manager.', False),
 ('  P3  Founder or co-founder at a smaller startup.', False),
 ('  P4  Generic company alias such as hr@ or info@ - lowest response rate.', False),
 ('', False),
 ('Rules:', True),
 ('  Every address here came from a real message in Aditya\'s Gmail. None are guessed.', False),
 ('  Never add an address that has not been verified - a bounce is a spam signal.', False),
 ('  Max 20 new outreach emails per day.', False),
 ('  Never email the same person twice for the same role without a reply.', False),
 ('  Stop immediately if someone asks not to be contacted, and add them to Do Not Contact.', False),
]
for i,(t,bold) in enumerate(LINES, start=1):
    c = g.cell(i,1,t)
    c.font = Font(name=F, size=12 if i==1 else 10, bold=bold, color=NAVY if i==1 else '000000')

wb.save('/home/user/AdiX/outreach/Outreach_Tracker.xlsx')
print('saved. rows:', len(ROWS))
